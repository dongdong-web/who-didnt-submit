"""从文件内容里挖署名 —— 文件名认不出时的最后一道防线。

只做「抽文本」这一件事，怎么用交给 matcher。

office 文件（docx / pptx）**不走 python-docx / python-pptx**：
它们本质就是 zip，正文分别在 `word/document.xml` 和 `ppt/slides/slideN.xml` 里。
直接抽 `<w:t>` / `<a:t>` 文本就够了 —— 省掉 lxml 和 Pillow 那几十兆依赖，
打包体积也不用为读内容再涨一截。

设计原则：**任何函数都不许抛异常**。一个损坏的 docx 不能让整次核对崩掉，
读不出来就返回空串，那个文件老老实实留在「认不出」里。
"""

from __future__ import annotations

import io
import re
import zipfile
from pathlib import Path

LIMIT = 20000          # 单个文件最多看这么多字
MAX_ZIP_MEMBERS = 30   # 压缩包最多翻这么多个成员

_TEXT_EXT = (".txt", ".md", ".csv", ".log", ".json", ".xml", ".html")

# docx 的 <w:t> 和 pptx 的 <a:t> 都表示「一段文本运行」
_RUN_RE = re.compile(r"<(?:w|a):t[^>]*>([^<]*)</(?:w|a):t>")
_DOCX_PART = re.compile(r"word/(?:document|header\d*|footer\d*)\.xml")
_PPTX_PART = re.compile(r"ppt/(?:slides/slide\d+|notesSlides/notesSlide\d+)\.xml")

_ENTITIES = (("&amp;", "&"), ("&lt;", "<"), ("&gt;", ">"),
             ("&quot;", '"'), ("&apos;", "'"))


def extract_text(path: Path) -> str:
    """尽力把文件里的文字抽出来。抽不到就返回空串。"""
    ext = path.suffix.lower()
    try:
        if ext == ".docx":
            return _office(path, _DOCX_PART)
        if ext == ".pptx":
            return _office(path, _PPTX_PART)
        if ext in (".xlsx", ".xlsm"):
            return _xlsx(path)
        if ext in _TEXT_EXT:
            return path.read_text(encoding="utf-8", errors="ignore")[:LIMIT]
        if ext == ".zip":
            return _zip(path)
    except Exception:
        return ""
    return ""


# --------------------------------------------------------------------------
# 各格式
# --------------------------------------------------------------------------


def _office(src, part_pattern: re.Pattern) -> str:
    """docx / pptx 都是 zip，里面是一堆 xml。src 可以是路径或 file-like。"""
    parts: list[str] = []
    total = 0
    with zipfile.ZipFile(src) as z:
        for name in z.namelist():
            if not part_pattern.fullmatch(name):
                continue
            try:
                xml = z.read(name).decode("utf-8", "ignore")
            except Exception:
                continue
            # 关键：run 之间**不加分隔符**。Word 会把「张伟」拆成两个 run，
            # 加了分隔符就再也拼不回一个完整名字了。
            text = "".join(_RUN_RE.findall(xml))
            for a, b in _ENTITIES:
                text = text.replace(a, b)
            parts.append(text)
            total += len(text)
            if total > LIMIT:
                break
    return "\n".join(parts)[:LIMIT]


def _xlsx(src) -> str:
    """表格用 openpyxl（它本来就要打包，读名单用的就是它）。"""
    from openpyxl import load_workbook

    wb = load_workbook(src, data_only=True, read_only=True)
    parts: list[str] = []
    total = 0
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            for c in row:
                if c is not None:
                    s = str(c)
                    parts.append(s)
                    total += len(s)
            if total > LIMIT:
                break
        if total > LIMIT:
            break
    wb.close()
    return "\n".join(parts)[:LIMIT]


def _zip(path: Path) -> str:
    """压缩包：成员文件名本身也是线索（`张三/实验一.docx`）。"""
    parts: list[str] = []
    with zipfile.ZipFile(path) as z:
        for info in z.infolist()[:MAX_ZIP_MEMBERS]:
            if info.is_dir():
                continue
            parts.append(info.filename)
            low = info.filename.lower()
            try:
                if low.endswith(_TEXT_EXT):
                    parts.append(z.read(info).decode("utf-8", "ignore"))
                elif low.endswith(".docx"):
                    parts.append(_office(io.BytesIO(z.read(info)), _DOCX_PART))
                elif low.endswith(".pptx"):
                    parts.append(_office(io.BytesIO(z.read(info)), _PPTX_PART))
                elif low.endswith((".xlsx", ".xlsm")):
                    parts.append(_xlsx(io.BytesIO(z.read(info))))
            except Exception:
                continue
    return "\n".join(parts)[:LIMIT]
